"""导出 Celery 任务 — 查询已归档用例并生成 Markdown 或 Excel 文件上传至 MinIO"""

from __future__ import annotations

import io
import logging
import re
from datetime import datetime, timezone
from uuid import UUID

from sqlalchemy import select, update

from src.platform_api.core.celery_app import celery_app
from src.platform_api.core.minio_client import ensure_bucket_exists, minio_client
from src.platform_api.core.settings import settings
from src.platform_api.models.enums import BatchStatus, ReviewStatus
from src.platform_api.models.knowledge import Document
from src.platform_api.models.public import System
from src.platform_api.models.testcase import ExportTask, TestBatch, TestCase
from src.testcase_generator.services.module_tree_classifier import classify_case_for_audit

logger = logging.getLogger(__name__)


def _split_cases(cases: list) -> tuple[list, list, list]:
    """按 verify 分桶（bucket）三路分流，收件人各不相同：
    - main：grounded（有据）/ unverified（未验上）→ 主用例集（可执行，交 QA）
    - needs_spec：undefined（PRD 未定义）/ ungrounded（无据编造）→ 需求澄清清单（交 PM 澄清）
    - to_fix：conflict（与 PRD 明文冲突的错误用例）→ 待修正用例（交测试/AI 改）
    旧批次 bucket 为 None → 归主集（行为不变）。"""
    main, clarification, to_fix = [], [], []
    for c in cases:
        bucket = getattr(c, "bucket", None)
        if bucket == "needs_spec":
            clarification.append(c)
        elif bucket == "to_fix":
            to_fix.append(c)
        else:
            main.append(c)
    return main, clarification, to_fix


@celery_app.task(name="platform_api.export", bind=True)
def export_task(
    self,
    export_id: str,
    scope: str,
    batch_id: str | None,
    system_id: str | None,
    format: str,
) -> dict:
    """
    导出任务执行流程：
    1. 查询已归档用例
    2. 生成 Markdown 或 Excel
    3. 上传 MinIO
    4. 更新 export_tasks.file_url + status=completed
    """
    import asyncio

    return asyncio.run(
        _execute_export(
            export_id=export_id,
            scope=scope,
            batch_id=batch_id,
            system_id=system_id,
            format=format,
        )
    )


async def _execute_export(
    export_id: str,
    scope: str,
    batch_id: str | None,
    system_id: str | None,
    format: str,
) -> dict:
    """内部异步执行逻辑"""
    from src.platform_api.core.database import get_session_factory

    session_factory = get_session_factory()

    try:
        async with session_factory() as session:
            # 1. 查询用例
            filters = []
            if scope == "batch" and batch_id:
                filters.append(TestCase.batch_id == UUID(batch_id))
            elif scope == "system" and system_id:
                # 查询系统下所有已归档批次的用例
                batch_ids_stmt = select(TestBatch.id).where(
                    TestBatch.system_id == UUID(system_id),
                    TestBatch.status == BatchStatus.ARCHIVED,
                )
                batch_ids_result = await session.execute(batch_ids_stmt)
                archived_batch_ids = [row[0] for row in batch_ids_result.fetchall()]
                if archived_batch_ids:
                    filters.append(TestCase.batch_id.in_(archived_batch_ids))
                else:
                    # 没有已归档批次，导出空文件
                    filters.append(TestCase.batch_id == UUID("00000000-0000-0000-0000-000000000000"))

            # 排除已删除的用例
            filters.append(TestCase.review_status != ReviewStatus.DELETED)

            stmt = (
                select(
                    TestCase,
                    System.name.label("system_name"),
                    Document.title.label("document_title"),
                )
                .join(TestBatch, TestCase.batch_id == TestBatch.id)
                .join(Document, TestBatch.document_id == Document.id)
                .join(System, TestBatch.system_id == System.id)
                .where(*filters)
                .order_by(TestCase.created_at.asc())
            )
            result = await session.execute(stmt)
            cases = []
            for case, system_name, document_title in result.all():
                setattr(case, "_export_system_name", system_name)
                setattr(case, "_export_document_title", document_title)
                cases.append(case)

            # 2. 三路分流：needs_spec → 需求澄清清单（交 PM）；to_fix → 待修正用例（交测试）；其余 → 主集
            main_cases, clarification_cases, fix_cases = _split_cases(cases)

            # 3. 生成文件内容
            if format == "markdown":
                content, content_type, file_ext = (
                    _generate_markdown(main_cases, clarification_cases, fix_cases),
                    "text/markdown",
                    "md",
                )
            else:
                content, content_type, file_ext = (
                    _generate_excel(main_cases, clarification_cases, fix_cases),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "xlsx",
                )

            # 4. 上传 MinIO
            ensure_bucket_exists()
            object_name = f"exports/{export_id}.{file_ext}"
            file_data = content if isinstance(content, bytes) else content.encode("utf-8")

            minio_client.put_object(
                bucket_name=settings.minio_bucket,
                object_name=object_name,
                data=io.BytesIO(file_data),
                length=len(file_data),
                content_type=content_type,
            )

            file_url = f"/{settings.minio_bucket}/{object_name}"

            # 5. 更新 export_tasks 记录
            update_stmt = (
                update(ExportTask)
                .where(ExportTask.id == UUID(export_id))
                .values(
                    status="completed",
                    file_url=file_url,
                    total_cases=len(main_cases),
                    completed_at=datetime.now(timezone.utc),
                )
            )
            await session.execute(update_stmt)
            await session.commit()

        logger.info(
            "Export %s completed: %d main + %d clarification + %d to_fix cases, format=%s",
            export_id,
            len(main_cases),
            len(clarification_cases),
            len(fix_cases),
            format,
        )
        # 注：total_cases 落库 ExportTask.total_cases（=主集可执行数）；下面两个计数仅随
        # celery result 返回 + 已写日志，不落库（ExportTask 无对应列），前端查 DB 暂取不到。
        return {
            "status": "completed",
            "export_id": export_id,
            "total_cases": len(main_cases),
            "clarification_cases": len(clarification_cases),
            "to_fix_cases": len(fix_cases),
        }

    except Exception as e:
        error_msg = f"{type(e).__name__}: {str(e)}"
        logger.exception("Export %s failed", export_id)

        # 更新失败状态
        async with session_factory() as session:
            update_stmt = (
                update(ExportTask)
                .where(ExportTask.id == UUID(export_id))
                .values(
                    status="failed",
                    error_message=error_msg,
                    completed_at=datetime.now(timezone.utc),
                )
            )
            await session.execute(update_stmt)
            await session.commit()

        return {"status": "failed", "export_id": export_id, "error": error_msg}


def _md_cell(value) -> str:
    """Markdown 表格单元格转义：竖线/换行会破坏表格结构，需替换。"""
    return (
        str(value if value is not None else "")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\n", "<br>")
        .replace("\r", "<br>")
    )


# Excel(openpyxl) 不接受的控制字符（除 \t\n\r），写入会抛 IllegalCharacterError
_XLSX_ILLEGAL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_TAPD_DIRECTORY_SEP_RE = re.compile(r"\s*(?:/|\\|>|＞|\||｜|»|→|—|–)\s*")
_TAPD_DIRECTORY_IGNORED_PARTS = {"_review_required", "unresolved_module", "待分类", "未匹配模块"}

TAPD_HEADERS = [
    "用例目录",
    "用例名称",
    "需求ID",
    "前置条件",
    "用例步骤",
    "预期结果",
    "用例类型",
    "用例状态",
    "用例等级",
    "创建人",
    "自测人",
]

TAPD_INSTRUCTIONS = [
    "[字段填写说明]“用例目录”请填写完整路径，用“-”分隔。如果目录为空，默认导入为“未规划目录”中；如果用例目录不存在，请在预览页面选择是否要自动创建目录。",
    "“用例名称”为必填项。",
    "“需求ID”请填写需求ID，多个需求ID以英文;号隔开。需求必须是本项目下的需求。",
    "“前置条件”请填写合法文本。",
    "“用例步骤”请填写合法文本。",
    "“预期结果”请填写合法文本。",
    "“用例类型”请填写：功能测试、性能测试、安全性测试、其他。",
    "“用例状态”请填写：正常、待更新、已废弃。",
    "“用例等级”请填写：高、中、低。",
    "支持多个人员传入,使用';'隔开 如: \"xxx;xxx\"",
    "单选人名字段，该字段将仅能填入1个成员。如若写入多个成员，只取第一个",
]

TAPD_PRIORITY_MAP = {
    "P0": "高",
    "P1": "中",
    "P2": "低",
    "P3": "低",
}


def _xlsx_safe(value) -> str:
    """清洗 Excel 不接受的控制字符，避免 openpyxl 抛 IllegalCharacterError。"""
    return _XLSX_ILLEGAL_RE.sub("", str(value if value is not None else ""))


def _case_provenance(case) -> dict:
    provenance = getattr(case, "provenance", None)
    return provenance if isinstance(provenance, dict) else {}


def _tapd_directory_parts(*values) -> list[str]:
    parts: list[str] = []
    for value in values:
        if value is None:
            continue
        normalized = _TAPD_DIRECTORY_SEP_RE.sub("-", str(value))
        for item in normalized.split("-"):
            part = item.strip(" -\t\r\n")
            if not part or part in _TAPD_DIRECTORY_IGNORED_PARTS:
                continue
            if parts and parts[-1] == part:
                continue
            parts.append(part)
    return parts


def _legacy_tapd_directory(case) -> str:
    provenance = _case_provenance(case)
    source = provenance.get("source_section")
    if not source:
        derived_from = provenance.get("derived_from")
        if isinstance(derived_from, list):
            source = "-".join(str(item) for item in derived_from if str(item).strip())
        elif derived_from:
            source = str(derived_from)
    if not source:
        return ""
    return "-".join(_tapd_directory_parts(source))


def _tapd_directory(case) -> str:
    """TAPD 用例目录：复用平台用例树坐标，按“系统-文档-模块-分支”输出完整路径。"""
    provenance = _case_provenance(case)
    classification = classify_case_for_audit(
        {
            "title": getattr(case, "title", ""),
            "provenance": provenance,
        }
    )
    module_name = classification.get("business_module")
    branch_path = classification.get("branch_path") or []

    if module_name not in _TAPD_DIRECTORY_IGNORED_PARTS:
        tree_parts = _tapd_directory_parts(
            getattr(case, "_export_system_name", None),
            getattr(case, "_export_document_title", None),
            module_name,
            *branch_path,
        )
        if tree_parts:
            return "-".join(tree_parts)

    legacy_directory = _legacy_tapd_directory(case)
    if legacy_directory:
        return legacy_directory

    return "-".join(
        _tapd_directory_parts(
            getattr(case, "_export_system_name", None),
            getattr(case, "_export_document_title", None),
        )
    )


def _tapd_priority(priority) -> str:
    """系统优先级到 TAPD 用例等级：P0=高、P1=中、P2/P3=低。"""
    return TAPD_PRIORITY_MAP.get(str(priority or "").upper(), "低")


def _format_tapd_steps(steps) -> str:
    """TAPD 步骤列只放操作与输入；逐步预期放到“预期结果”列。"""
    if not isinstance(steps, list):
        return _format_field(steps)

    lines: list[str] = []
    for i, step in enumerate(steps, 1):
        if isinstance(step, dict):
            action = step.get("action") or step.get("step") or ""
            input_data = step.get("input_data")
            line = str(action or step)
            if input_data:
                line = f"{line}（输入：{input_data}）"
        else:
            line = str(step)
        lines.append(f"{i}. {line}")
    return "\n".join(lines)


def _format_tapd_expected(case) -> str:
    steps = getattr(case, "steps", None)
    if isinstance(steps, list):
        expected_lines = [
            f"{i}. {step.get('expected_result')}"
            for i, step in enumerate(steps, 1)
            if isinstance(step, dict) and step.get("expected_result")
        ]
        if expected_lines:
            return "\n".join(expected_lines)
    return _format_field(getattr(case, "expected_results", None))


def _tapd_row(case) -> list:
    return [
        _xlsx_safe(_tapd_directory(case)),
        _xlsx_safe(case.title),
        "",
        _xlsx_safe(_format_field(case.preconditions)),
        _xlsx_safe(_format_tapd_steps(case.steps)),
        _xlsx_safe(_format_tapd_expected(case)),
        "",
        "",
        _xlsx_safe(_tapd_priority(case.priority)),
        "",
        "",
    ]


def _review_row(c) -> tuple[str, str, str]:
    """从用例提取 (verdict, rationale, prd_evidence)，供需求澄清/待修正清单复用。

    export_task 读取的是 ORM TestCase：verdict 为独立列、rationale/prd_evidence
    在 verification(JSONB) 内；旧数据 verification 为 None 时全部回退空串。
    """
    v = getattr(c, "verification", None)
    if isinstance(v, dict):
        rationale = v.get("rationale", "")
        evidence = v.get("prd_evidence", "")
    else:
        rationale = ""
        evidence = ""
    verdict = getattr(c, "verdict", "") or ""
    return verdict, rationale, evidence


def _append_md_review_section(lines: list[str], title: str, rows: list) -> None:
    """向 markdown 追加一个核验清单段（需求澄清清单 / 待修正用例 共用同结构）。"""
    lines.append(f"## {title}\n")
    lines.append("| 序号 | 标题 | verdict | 判定理由 | PRD依据 |\n")
    lines.append("| --- | --- | --- | --- | --- |\n")
    for i, c in enumerate(rows, 1):
        verdict, rationale, evidence = _review_row(c)
        lines.append(
            f"| {i} | {_md_cell(c.title)} | {_md_cell(verdict)} | {_md_cell(rationale)} | {_md_cell(evidence)} |\n"
        )


def _append_excel_review_sheet(wb, title: str, rows: list) -> None:
    """向 workbook 追加一个核验清单 sheet（需求澄清清单 / 待修正用例 共用同结构）。"""
    ws = wb.create_sheet(title)
    ws.append(["序号", "标题", "verdict", "判定理由", "PRD依据"])
    for i, c in enumerate(rows, 1):
        verdict, rationale, evidence = _review_row(c)
        ws.append([i, _xlsx_safe(c.title), _xlsx_safe(verdict), _xlsx_safe(rationale), _xlsx_safe(evidence)])


def _generate_markdown(cases: list, clarification: list | None = None, to_fix: list | None = None) -> str:
    """生成 Markdown 评审视图：主表按 TAPD 字段口径，问题用例保留为独立清单。"""
    lines: list[str] = ["| " + " | ".join(_md_cell(header) for header in TAPD_HEADERS) + " |\n"]
    lines.append("| " + " | ".join("---" for _ in TAPD_HEADERS) + " |\n")
    for case in cases:
        lines.append("| " + " | ".join(_md_cell(value) for value in _tapd_row(case)) + " |\n")
    if clarification:
        lines.append("\n")
        _append_md_review_section(lines, "需求澄清清单（待 PM 确认，未计入可执行用例）", clarification)
    if to_fix:
        lines.append("\n")
        _append_md_review_section(lines, "待修正用例（与 PRD 冲突，需测试/AI 修正，未计入可执行用例）", to_fix)
    return "".join(lines)


def _generate_excel(cases: list, clarification: list | None = None, to_fix: list | None = None) -> bytes:
    """将可执行用例生成 TAPD 导入模板 Excel。

    clarification / to_fix 不进入 TAPD 导入文件：它们不是可执行用例，需留在 Markdown
    评审文档中处理，避免外部测试管理工具误导入。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        logger.info("openpyxl 不可用，降级为 TAPD CSV 导出")
        return _generate_csv_fallback(cases, clarification, to_fix)

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    ws.append(TAPD_HEADERS)
    ws.append(TAPD_INSTRUCTIONS)

    header_fill = PatternFill(fill_type="solid", fgColor="C4BD97")
    instruction_fill = PatternFill(fill_type="solid", fgColor="EBF1DE")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(name="Calibri", size=16, bold=True, color="000000")
        cell.alignment = Alignment(vertical="center")
    for cell in ws[2]:
        cell.fill = instruction_fill
        cell.font = Font(name="Calibri", size=11, color="000000")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

    for case in cases:
        ws.append(_tapd_row(case))

    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 24,
        "B": 32,
        "C": 13,
        "D": 28,
        "E": 42,
        "F": 42,
        "G": 13,
        "H": 13,
        "I": 13,
        "J": 13,
        "K": 13,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A3"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def _generate_csv_fallback(cases: list, clarification: list | None = None, to_fix: list | None = None) -> bytes:
    """降级 CSV 导出（openpyxl 不可用时），保持 TAPD 模板列顺序。"""
    import csv

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TAPD_HEADERS)
    writer.writerow(TAPD_INSTRUCTIONS)

    for case in cases:
        writer.writerow(_tapd_row(case))

    return output.getvalue().encode("utf-8-sig")


def _format_field(value) -> str:
    """将 JSONB 字段格式化为可读字符串"""
    if isinstance(value, list):
        items = []
        for item in value:
            if isinstance(item, dict):
                items.append("; ".join(f"{k}={v}" for k, v in item.items()))
            else:
                items.append(str(item))
        return "\n".join(items)
    elif isinstance(value, dict):
        return "; ".join(f"{k}={v}" for k, v in value.items())
    return str(value) if value else ""
